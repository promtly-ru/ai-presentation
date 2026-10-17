"""Собирает файлы для скачивания из index.html через headless Microsoft Edge.

Результат:
  downloads/promptly-ai-presentation.pdf   — печать страницы (@page 1920x1080, по слайду на страницу)
  downloads/promptly-ai-presentation.pptx  — слайды-изображения 16:9
  img/slide-N.png                          — кадры слайдов
  img/preview-cover.jpg, img/preview-calc.jpg — превью 1280x800 для сайта

Запуск: python tools/build_downloads.py
"""
import subprocess
import sys
import tempfile
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.util import Emu

ROOT = Path(__file__).resolve().parent.parent
INDEX = ROOT / 'index.html'
DOWNLOADS = ROOT / 'downloads'
IMG = ROOT / 'img'
SLIDES = 7
EDGE_CANDIDATES = [
    Path(r'C:\Program Files (x86)\Microsoft\Edge\Application\msedge.exe'),
    Path(r'C:\Program Files\Microsoft\Edge\Application\msedge.exe'),
]


def edge() -> str:
    for p in EDGE_CANDIDATES:
        if p.exists():
            return str(p)
    sys.exit('Microsoft Edge не найден')


def run_edge(args: list[str], profile: str) -> None:
    cmd = [edge(), '--headless=new', '--disable-gpu', '--hide-scrollbars', '--no-first-run',
           f'--user-data-dir={profile}', '--virtual-time-budget=9000', *args]
    subprocess.run(cmd, check=True, timeout=120, capture_output=True)


def main() -> None:
    DOWNLOADS.mkdir(exist_ok=True)
    IMG.mkdir(exist_ok=True)
    url = INDEX.as_uri()
    profile = tempfile.mkdtemp(prefix='edge-deck-')

    shots = []
    for i in range(1, SLIDES + 1):
        out = IMG / f'slide-{i}.png'
        run_edge(['--window-size=1920,1080', f'--screenshot={out}', f'{url}?static#{i}'], profile)
        with Image.open(out) as im:
            if im.size != (1920, 1080):
                sys.exit(f'{out.name}: размер {im.size}, ожидался 1920x1080')
        shots.append(out)
        print('кадр', out.name)

    pdf = DOWNLOADS / 'promptly-ai-presentation.pdf'
    run_edge(['--no-pdf-header-footer', f'--print-to-pdf={pdf}', f'{url}?static'], profile)
    print('pdf', pdf.name, pdf.stat().st_size // 1024, 'КБ')

    prs = Presentation()
    prs.slide_width = Emu(12192000)   # 13.333 дюйма, 16:9
    prs.slide_height = Emu(6858000)   # 7.5 дюйма
    blank = prs.slide_layouts[6]
    for shot in shots:
        s = prs.slides.add_slide(blank)
        s.shapes.add_picture(str(shot), 0, 0, width=prs.slide_width, height=prs.slide_height)
    pptx = DOWNLOADS / 'promptly-ai-presentation.pptx'
    prs.save(pptx)
    print('pptx', pptx.name, pptx.stat().st_size // 1024, 'КБ')

    # Превью 16:10 для карточки на сайте: масштаб по высоте и обрезка по бокам (поля слайда это позволяют)
    for n, name in ((1, 'preview-cover.jpg'), (5, 'preview-calc.jpg')):
        with Image.open(IMG / f'slide-{n}.png') as im:
            im = im.convert('RGB').resize((1422, 800), Image.LANCZOS)
            left = (1422 - 1280) // 2
            im.crop((left, 0, left + 1280, 800)).save(IMG / name, quality=86, optimize=True)
        print('превью', name)


if __name__ == '__main__':
    main()
